"""Small, local presentation and chart tools. Files keep the run's access boundary."""

from __future__ import annotations

import asyncio
import hashlib
import json
import re
from html import escape
from io import BytesIO
from typing import Annotated
from uuid import uuid4

from langchain_core.tools import tool
from pydantic import BaseModel, Field

from app.agent.models import AgentFile
from app.core.config import settings


class PresentationSlide(BaseModel):
    title: str = Field(min_length=1, max_length=100)
    points: list[Annotated[str, Field(min_length=1, max_length=140)]] = Field(
        default_factory=list, max_length=5
    )
    notes: str = Field(default="", max_length=3000)


class ChartValue(BaseModel):
    label: str = Field(min_length=1, max_length=40)
    value: float = Field(ge=0, le=1e12, allow_inf_nan=False)


def presentation_bytes(title, slides, author):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title, deck.core_properties.author = title, author

    def text(slide, value, x, y, width, height, size, color, bold=False):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap, frame.vertical_anchor = True, MSO_ANCHOR.MIDDLE
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        paragraph = frame.paragraphs[0]
        paragraph.text = value
        paragraph.font.name, paragraph.font.size = "Arial", Pt(size)
        paragraph.font.bold, paragraph.font.color.rgb = bold, RGBColor.from_string(color)

    for index, content in enumerate(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        dark = index == 0
        background, foreground = ("183C32", "FFFFFF") if dark else ("FFFFFF", "183C32")
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(background)
        text(slide, content.title, 0.65, 0.5, 12, 1.2, 34, foreground, True)
        for row, point in enumerate(content.points):
            y = 2.05 + row * 0.88
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(y), Inches(12), Inches(0.76)
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string("264D41" if dark else "EDF3F0")
            shape.line.fill.background()
            text(slide, str(row + 1).zfill(2), 0.9, y + 0.06, 0.55, 0.64, 20, "78B99E", True)
            text(slide, point, 1.7, y + 0.06, 10.65, 0.64, 18, foreground)
        text(slide, author, 0.65, 7.03, 10, 0.22, 10, "9DB9AD" if dark else "60746B")
        text(slide, f"{index + 1:02}", 12.1, 7.03, 0.55, 0.22, 10, "9DB9AD" if dark else "60746B")
        if content.notes:
            slide.notes_slide.notes_text_frame.text = content.notes
    stream = BytesIO()
    deck.save(stream)
    content = stream.getvalue()
    # Read back the file before claiming successful generation.
    assert len(Presentation(BytesIO(content)).slides) == len(slides)
    return content


def chart_bytes(title, values):
    height = 145 + len(values) * 48
    maximum = max((point.value for point in values), default=0) or 1
    parts = [
        '<svg xmlns="http://www.w3.org/2000/svg" width="1000" '
        f'height="{height}" viewBox="0 0 1000 {height}" role="img">',
        f"<title>{escape(title)}</title>",
        '<rect width="100%" height="100%" fill="white"/>',
        '<g font-family="Arial, sans-serif" fill="#183C32">'
        f'<text x="35" y="55" font-size="27" font-weight="bold">{escape(title)}</text>',
    ]
    for index, point in enumerate(values):
        y, width = 95 + index * 48, point.value / maximum * 510
        parts.extend(
            [
                f'<text x="35" y="{y + 22}" font-size="16">{escape(point.label)}</text>',
                f'<rect x="360" y="{y}" width="{width:.2f}" height="30" rx="4" fill="#35755C"/>',
                f'<text x="{382 + width:.2f}" y="{y + 22}" font-size="16">{point.value:g}</text>',
            ]
        )
    return ("".join(parts) + "</g></svg>").encode()


async def save_file(db, run, title, extension, media_type, content):
    if len(content) > 5 * 1024 * 1024:
        raise ValueError("Generated file exceeds the 5 MB limit.")
    filename = (re.sub(r"[^a-zA-Z0-9_-]+", "-", title).strip("-")[:100] or "document") + extension
    row = AgentFile(
        id=uuid4(),
        run_id=run.id,
        filename=filename,
        media_type=media_type,
        content=content,
        sha256=hashlib.sha256(content).hexdigest(),
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    if hashlib.sha256(row.content).hexdigest() != row.sha256:
        raise ValueError("Generated file failed readback verification.")
    return json.dumps(
        {
            "file": {
                "id": str(row.id),
                "filename": row.filename,
                "bytes": len(content),
                "sha256": row.sha256,
                "url": f"{settings.frontend_url.rstrip('/')}/files/{row.id}",
                "access": "Sign in to the company account to download.",
            }
        }
    )


def file_tools(db, employee, run, slug):
    @tool
    async def create_presentation(
        title: Annotated[str, Field(min_length=1, max_length=100)],
        slides: Annotated[list[PresentationSlide], Field(min_length=1, max_length=12)],
    ) -> str:
        """Create an editable PowerPoint (.pptx), save it privately, and return its download URL.

        Pass one item per requested slide. Each has a title, up to 5 short points (140 characters
        each), and optional speaker notes. Use confirmed facts or label illustrative data.
        Include the returned download URL in your reply; downloading requires company sign-in.
        """
        content = await asyncio.to_thread(presentation_bytes, title, slides, employee.name)
        return await save_file(
            db,
            run,
            title,
            ".pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            content,
        )

    @tool
    async def create_chart(
        title: Annotated[str, Field(min_length=1, max_length=80)],
        values: Annotated[list[ChartValue], Field(min_length=1, max_length=12)],
    ) -> str:
        """Create a downloadable SVG bar chart from explicit labels and nonnegative values.

        Do not invent company metrics. Return the download URL; company sign-in is required.
        """
        return await save_file(db, run, title, ".svg", "image/svg+xml", chart_bytes(title, values))

    return [create_presentation] if slug == "pitchdeck" else [create_chart]
