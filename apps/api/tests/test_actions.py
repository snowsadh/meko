"""Action wiring checks with real storage and explicit remote-service/model stubs."""

import json
from importlib import import_module
from uuid import UUID

import pytest
from langchain_core.messages import AIMessage
from langchain_core.tools import tool

from app.agent.build import run_agent
from app.agent.tools import executor
from app.employees.models import Employee
from tests.conftest import setup_company

connections = import_module("app.mcp.router")


@pytest.mark.asyncio
async def test_mcp_connection_discovery_employee_scope_and_failure(product, monkeypatch):
    client, _, _ = product
    headers, org, support, hr = await setup_company(client)
    base = f"/api/organizations/{org['id']}/employees/{support['id']}/mcp-connections"
    calls = []

    @tool
    async def web_search_web_search_exa(query: str) -> str:
        """Search public pages."""
        calls.append(query)
        return json.dumps({"url": "https://example.com/support", "text": "Public support docs"})

    async def discover(slug, **kwargs):
        if slug == "n8n":
            raise ValueError("Service unavailable")
        return [web_search_web_search_exa]

    monkeypatch.setattr(connections, "remote_tools", discover)
    monkeypatch.setattr(executor, "remote_tools", discover)
    response = await client.post(base + "/web_search", headers=headers, json={"credential": ""})
    assert response.status_code == 201, response.text
    assert response.json()["auth_type"] == "none"
    response = await client.post(
        base + "/n8n",
        headers=headers,
        json={
            "credential": "synthetic-token",
            "server_url": "https://example.com/mcp",
        },
    )
    assert response.status_code == 502
    assert len((await client.get(base, headers=headers)).json()["connections"]) == 1
    assert (await client.post(base + "/n8n", headers=headers, json={})).status_code == 422

    class SearchModel:
        def bind_tools(self, tools):
            assert "web_search_web_search_exa" in {item.name for item in tools}
            return self

        async def ainvoke(self, messages):
            if not calls:
                return AIMessage(
                    content="",
                    tool_calls=[
                        {
                            "name": "web_search_web_search_exa",
                            "args": {"query": "public support docs"},
                            "id": "search-1",
                        }
                    ],
                )
            return AIMessage(content="Found public docs: https://example.com/support")

    from app.core.database import get_db
    from app.main import app

    async for db in app.dependency_overrides[get_db]():
        hr_employee = await db.get(Employee, UUID(hr["id"]))
        assert await executor.connected_tools(db, hr_employee) == []
        run = await run_agent(
            db,
            org_id=UUID(org["id"]),
            employee_id=UUID(support["id"]),
            user_id=UUID((await client.get("/api/auth/me", headers=headers)).json()["id"]),
            message="Find public support docs",
            use_meko=False,
            model=SearchModel(),
        )
        assert run.status == "completed"
        assert run.tools[0]["status"] == "completed"
        assert "https://example.com/support" in run.tools[0]["result"]
        await client.delete(base + "/web_search", headers=headers)
        assert await executor.connected_tools(db, await db.get(Employee, UUID(support["id"]))) == []
        # A stale/unavailable connection cannot break ordinary chat.
        await client.post(base + "/web_search", headers=headers, json={})

        async def unavailable(*args, **kwargs):
            raise TimeoutError()

        monkeypatch.setattr(executor, "remote_tools", unavailable)
        tools = await executor.connected_tools(db, await db.get(Employee, UUID(support["id"])), run)
        assert tools == []
        assert run.tools[-1]["status"] == "failed"
        break


@pytest.mark.asyncio
async def test_generated_files_real_bytes_download_ownership_and_disconnect(product):
    import hashlib
    from io import BytesIO

    from pptx import Presentation
    from sqlalchemy import func, select

    from app.agent.models import AgentFile
    from app.core.database import get_db
    from app.main import app

    client, fake, _ = product
    headers, org, support, _ = await setup_company(client)
    base = f"/api/organizations/{org['id']}/employees/{support['id']}/mcp-connections"
    for slug in ["pitchdeck", "visualization"]:
        response = await client.post(base + "/" + slug, headers=headers, json={})
        assert response.status_code == 201, response.text
    user_id = UUID((await client.get("/api/auth/me", headers=headers)).json()["id"])

    class FileModel:
        def bind_tools(self, tools):
            assert {tool.name for tool in tools} == {"create_presentation", "create_chart"}
            return self

        async def ainvoke(self, messages):
            if messages[-1].type == "human":
                return AIMessage(
                    content="",
                    tool_calls=[
                        {
                            "name": "create_presentation",
                            "id": "pptx-1",
                            "args": {
                                "title": "Support Handbook",
                                "slides": [
                                    {
                                        "title": title,
                                        "points": ["Synthetic test data"],
                                        "notes": "Speaker notes",
                                    }
                                    for title in ["Support Overview", "Triage", "Escalation"]
                                ],
                            },
                        },
                        {
                            "name": "create_chart",
                            "id": "chart-1",
                            "args": {
                                "title": "Synthetic tickets",
                                "values": [
                                    {"label": "<script>alert(1)</script>", "value": 12},
                                    {"label": "Resolved", "value": 8},
                                ],
                            },
                        },
                    ],
                )
            return AIMessage(content="Files created. Download using the returned URLs.")

    args = dict(
        org_id=UUID(org["id"]),
        employee_id=UUID(support["id"]),
        user_id=user_id,
        message="Create a deck and chart using synthetic data",
        use_meko=False,
        request_key="files-once",
        model=FileModel(),
    )
    async for db in app.dependency_overrides[get_db]():
        run = await run_agent(db, **args)
        assert run.status == "completed", run.error
        assert [event["status"] for event in run.tools] == ["completed", "completed"], run.tools
        files = [json.loads(event["result"])["file"] for event in run.tools]
        for file in files:
            response = await client.get("/api/agent/files/" + file["id"], headers=headers)
            assert response.status_code == 200, response.text
            assert hashlib.sha256(response.content).hexdigest() == file["sha256"]
            assert response.headers["x-content-sha256"] == file["sha256"]
            assert response.headers["cache-control"] == "private, no-store"
            if file["filename"].endswith(".pptx"):
                deck = Presentation(BytesIO(response.content))
                assert len(deck.slides) == 3
                assert "Speaker notes" in deck.slides[1].notes_slide.notes_text_frame.text
                assert "Triage" in "\n".join(
                    shape.text for shape in deck.slides[1].shapes if shape.has_text_frame
                )
            else:
                assert b"<script>" not in response.content
                assert b"&lt;script&gt;" in response.content
                assert b">12</text>" in response.content
        assert (await run_agent(db, **args)).id == run.id
        assert await db.scalar(select(func.count()).select_from(AgentFile)) == 2
        assert fake.calls == []
        # Another company owner and an unauthenticated caller cannot download the files.
        other, *_ = await setup_company(client, suffix="other-files")
        assert (
            await client.get("/api/agent/files/" + files[0]["id"], headers=other)
        ).status_code == 404
        assert (await client.get("/api/agent/files/" + files[0]["id"])).status_code == 401
        await client.delete(base + "/pitchdeck", headers=headers)
        employee = await db.get(Employee, UUID(support["id"]))
        assert [item.name for item in await executor.connected_tools(db, employee, run)] == [
            "create_chart"
        ]
        break
